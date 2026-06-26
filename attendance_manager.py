import os
import time
from dataclasses import dataclass
from datetime import date, datetime, time as dtime
from typing import Iterable

from openpyxl import Workbook, load_workbook

import config


IN_ACTION = "IN"
OUT_ACTION = "OUT"


@dataclass(frozen=True)
class AttendanceEvent:
    action: str
    timestamp: datetime

    @classmethod
    def from_row(cls, row) -> "AttendanceEvent":
        return cls(action=str(row.action or "").upper(), timestamp=row.ts)


@dataclass(frozen=True)
class AttendanceDecision:
    accepted: bool
    action: str | None
    timestamp: datetime
    level: str
    message: str
    reason: str = ""
    remaining_seconds: int = 0


@dataclass(frozen=True)
class DailyAttendanceSummary:
    user_id: int | None
    name: str
    day: date
    earliest_in: datetime | None
    latest_out: datetime | None
    status: str
    late: bool
    early: bool
    worked_minutes: int
    events_count: int

    def to_dict(self) -> dict:
        return {
            "date": self.day.isoformat(),
            "user_id": self.user_id,
            "name": self.name,
            "earliest_in": self.earliest_in.isoformat() if self.earliest_in else None,
            "latest_out": self.latest_out.isoformat() if self.latest_out else None,
            "status": self.status,
            "late": self.late,
            "early": self.early,
            "worked_minutes": self.worked_minutes,
            "events_count": self.events_count,
        }


@dataclass
class AttendancePolicy:
    shift_start: dtime
    shift_end: dtime
    min_in_to_out: int
    min_out_to_in: int
    trigger_cooldown: int


class AttendanceManager:
    """Owns attendance business rules and the operator-facing event message."""

    def __init__(self, file_path: str | None = None, policy: AttendancePolicy | None = None):
        self.file_path = file_path or getattr(config, "ATTENDANCE_XLSX", "attendance.xlsx")
        self.policy = policy or AttendancePolicy(
            shift_start=self._parse_time(getattr(config, "SHIFT_START_TIME", "09:00")),
            shift_end=self._parse_time(getattr(config, "SHIFT_END_TIME", "17:00")),
            min_in_to_out=getattr(config, "ATT_MIN_IN_TO_OUT_SECONDS", 60),
            min_out_to_in=getattr(config, "ATT_MIN_OUT_TO_IN_SECONDS", 60),
            trigger_cooldown=getattr(config, "ATT_TRIGGER_COOLDOWN", 5),
        )
        self.last_attendance_trigger: dict[str, datetime] = {}
        self.last_event: dict | None = None
        self.last_event_ts = 0.0
        self._init_attendance_file()

    @property
    def min_in_to_out(self) -> int:
        return self.policy.min_in_to_out

    @min_in_to_out.setter
    def min_in_to_out(self, value: int) -> None:
        self.policy.min_in_to_out = int(value)

    @property
    def min_out_to_in(self) -> int:
        return self.policy.min_out_to_in

    @min_out_to_in.setter
    def min_out_to_in(self, value: int) -> None:
        self.policy.min_out_to_in = int(value)

    @property
    def shift_start(self) -> dtime:
        return self.policy.shift_start

    @shift_start.setter
    def shift_start(self, value: dtime) -> None:
        self.policy.shift_start = value

    @property
    def shift_end(self) -> dtime:
        return self.policy.shift_end

    @shift_end.setter
    def shift_end(self, value: dtime) -> None:
        self.policy.shift_end = value

    def configure(
        self,
        shift_start: str,
        shift_end: str,
        min_in_to_out: int,
        min_out_to_in: int,
        trigger_cooldown: int | None = None,
    ) -> None:
        self.policy.shift_start = self._parse_time(shift_start)
        self.policy.shift_end = self._parse_time(shift_end)
        self.policy.min_in_to_out = int(min_in_to_out)
        self.policy.min_out_to_in = int(min_out_to_in)
        if trigger_cooldown is not None:
            self.policy.trigger_cooldown = int(trigger_cooldown)

    def handle(
        self,
        name: str,
        events: Iterable[AttendanceEvent] | None = None,
        identity_key: str | int | None = None,
        now: datetime | None = None,
        mirror_to_xlsx: bool = True,
    ) -> AttendanceDecision:
        now = now or datetime.now()
        key = str(identity_key or name)
        records = list(events) if events is not None else self._get_today_records(name)
        decision = self.decide(name=name, events=records, identity_key=key, now=now)
        self.last_attendance_trigger[key] = now

        if decision.accepted and decision.action:
            if mirror_to_xlsx:
                self._log(name, decision.action, decision.timestamp)

        self._emit_decision(name, decision)
        return decision

    def decide(
        self,
        name: str,
        events: Iterable[AttendanceEvent],
        identity_key: str,
        now: datetime,
    ) -> AttendanceDecision:
        last_trigger = self.last_attendance_trigger.get(identity_key)
        if last_trigger:
            elapsed = (now - last_trigger).total_seconds()
            if elapsed < self.policy.trigger_cooldown:
                remaining = max(1, int(self.policy.trigger_cooldown - elapsed))
                return AttendanceDecision(
                    accepted=False,
                    action=None,
                    timestamp=now,
                    level="warning",
                    message=f"Recently processed. Please wait {remaining} seconds.",
                    reason="trigger_cooldown",
                    remaining_seconds=remaining,
                )

        records = sorted(events, key=lambda item: item.timestamp)
        if not records:
            return self._accepted(IN_ACTION, now)

        last = records[-1]
        last_action = (last.action or "").upper()
        elapsed = (now - last.timestamp).total_seconds()

        if last_action == IN_ACTION:
            if elapsed < self.policy.min_in_to_out:
                remaining = max(1, int(self.policy.min_in_to_out - elapsed))
                return AttendanceDecision(
                    accepted=False,
                    action=None,
                    timestamp=now,
                    level="warning",
                    message=f"Already checked in at {last.timestamp:%H:%M:%S}. Please wait {remaining} seconds to check out.",
                    reason="min_in_to_out",
                    remaining_seconds=remaining,
                )
            return self._accepted(OUT_ACTION, now)

        if last_action == OUT_ACTION:
            if elapsed < self.policy.min_out_to_in:
                remaining = max(1, int(self.policy.min_out_to_in - elapsed))
                return AttendanceDecision(
                    accepted=False,
                    action=None,
                    timestamp=now,
                    level="warning",
                    message=f"Already checked out at {last.timestamp:%H:%M:%S}. Please wait {remaining} seconds to check in.",
                    reason="min_out_to_in",
                    remaining_seconds=remaining,
                )
            return self._accepted(IN_ACTION, now)

        return AttendanceDecision(
            accepted=False,
            action=None,
            timestamp=now,
            level="error",
            message=f"Unsupported attendance action for {name}.",
            reason="invalid_previous_action",
        )

    def summarize_day(
        self,
        name: str,
        day: date,
        events: Iterable[AttendanceEvent] | None = None,
        user_id: int | None = None,
    ) -> DailyAttendanceSummary:
        records = list(events) if events is not None else self._get_day_records_from_xlsx(name, day)
        records = sorted(records, key=lambda item: item.timestamp)

        ins = [item.timestamp for item in records if item.action.upper() == IN_ACTION]
        outs = [item.timestamp for item in records if item.action.upper() == OUT_ACTION]

        earliest_in = min(ins) if ins else None
        latest_out = max(outs) if outs else None
        late = bool(earliest_in and earliest_in.time() > self.policy.shift_start)
        early = bool(latest_out and latest_out.time() < self.policy.shift_end)

        if not records:
            status = "absent"
        elif earliest_in and latest_out and not late and not early:
            status = "full"
        else:
            status = "incomplete"

        worked_minutes = 0
        if earliest_in and latest_out and latest_out > earliest_in:
            worked_minutes = int((latest_out - earliest_in).total_seconds() // 60)

        return DailyAttendanceSummary(
            user_id=user_id,
            name=name,
            day=day,
            earliest_in=earliest_in,
            latest_out=latest_out,
            status=status,
            late=late,
            early=early,
            worked_minutes=worked_minutes,
            events_count=len(records),
        )

    def load_rows(self) -> list[dict]:
        if not os.path.exists(self.file_path):
            return []

        wb = load_workbook(self.file_path)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            name, action, day, raw_time = row
            rows.append({
                "name": name,
                "action": action,
                "date": str(day),
                "time": str(raw_time),
            })
        return rows

    def _accepted(self, action: str, now: datetime) -> AttendanceDecision:
        label = "Check-in" if action == IN_ACTION else "Check-out"
        return AttendanceDecision(
            accepted=True,
            action=action,
            timestamp=now,
            level="success",
            message=f"{label} successful at {now:%H:%M:%S}.",
            reason="accepted",
        )

    def _emit_decision(self, name: str, decision: AttendanceDecision) -> None:
        if decision.accepted and decision.action:
            action_label = "Check-in" if decision.action == IN_ACTION else "Check-out"
            text = f"{name}\n{action_label} successful\nTime: {decision.timestamp:%H:%M:%S}"
        else:
            text = f"{name}\n{decision.message}"

        self.last_event = {
            "action": (decision.action or decision.reason or "notice").lower(),
            "level": decision.level,
            "text": text,
            "accepted": decision.accepted,
            "reason": decision.reason,
        }
        self.last_event_ts = time.time()

    def _init_attendance_file(self) -> None:
        if os.path.exists(self.file_path):
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(["name", "action", "date", "time"])
        wb.save(self.file_path)

    def _log(self, name: str, action: str, now: datetime) -> None:
        wb = load_workbook(self.file_path)
        ws = wb.active
        ws.append([name, action, now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S")])
        wb.save(self.file_path)

    def _get_today_records(self, name: str) -> list[AttendanceEvent]:
        return self._get_day_records_from_xlsx(name, date.today())

    def _get_day_records_from_xlsx(self, name: str, day: date) -> list[AttendanceEvent]:
        wanted = day.strftime("%Y-%m-%d")
        records = []
        for row in self.load_rows():
            if row["name"] != name or row["date"] != wanted:
                continue
            try:
                ts = datetime.strptime(f"{row['date']} {row['time']}", "%Y-%m-%d %H:%M:%S")
            except Exception:
                continue
            records.append(AttendanceEvent(action=str(row["action"]).upper(), timestamp=ts))
        return records

    @staticmethod
    def _parse_time(value: str) -> dtime:
        try:
            return datetime.strptime(value, "%H:%M").time()
        except Exception:
            return dtime(hour=9, minute=0)
